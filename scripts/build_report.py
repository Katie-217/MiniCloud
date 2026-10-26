"""Build the local report as HTML/PDF/DOCX from one Markdown source."""
import html
from pathlib import Path
import re
import markdown
from docx import Document
from docx.shared import Inches, Pt
from playwright.sync_api import sync_playwright
from browser import launch
from lab import ROOT
from progress import set_status

folder = ROOT / "report"
source = folder / "BaoCao_Local_521H0092.md"
text = source.read_text(encoding="utf-8")
with sync_playwright() as p:
    browser = launch(p)
    page = browser.new_page(viewport={"width": 1240, "height": 660})
    page.goto((folder / "architecture.svg").as_uri())
    page.locator("svg").screenshot(path=str(folder / "architecture.png"))
    body = markdown.markdown(text, extensions=["tables", "fenced_code"])
    style = '''body{font:11pt Arial,sans-serif;line-height:1.6;color:#19303d}h1{font-size:25pt;color:#125d55}h2{font-size:18pt;color:#125d55;border-bottom:1px solid #b8cecc;page-break-before:always}h3{font-size:13pt;page-break-after:avoid}table{width:100%;border-collapse:collapse;font-size:9pt}th,td{border:1px solid #b8c8cd;padding:7px;vertical-align:top}th{background:#edf4f4}tr{break-inside:avoid}img{max-width:100%;max-height:225mm;object-fit:contain;display:block;margin:16px auto}p:has(img){break-inside:avoid}pre{white-space:pre-wrap;word-break:break-word;background:#f1f5f6;padding:12px;font-size:9pt}a{color:#126b77;overflow-wrap:anywhere}code{font-size:9pt}'''
    report_html = '<!doctype html><html lang="vi"><head><meta charset="utf-8"><title>MyMiniCloud — Báo cáo local</title><style>' + style + '</style></head><body>' + body + '</body></html>'
    html_path = source.with_suffix(".html")
    html_path.write_text(report_html, encoding="utf-8")
    page.goto(html_path.as_uri())
    page.evaluate("document.fonts.ready")
    page.pdf(path=str(source.with_suffix(".pdf")), format="A4", print_background=True,
        margin={"top": "18mm", "bottom": "18mm", "left": "18mm", "right": "18mm"},
        display_header_footer=True, header_template="<span></span>",
        footer_template='<div style="font:9px Arial;width:100%;padding:0 18mm;color:#607080">MyMiniCloud · 521H0092 · LOCAL — AWS pending<span style="float:right"><span class="pageNumber"></span> / <span class="totalPages"></span></span></div>')
    browser.close()

document = Document()
normal = document.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10)
section = document.sections[0]
section.top_margin = Inches(.7)
section.bottom_margin = Inches(.7)
section.left_margin = Inches(.7)
section.right_margin = Inches(.7)

def plain(line):
    line = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", line)
    return line.replace("**", "").replace("`", "")

lines = text.splitlines()
index = 0
code = False
while index < len(lines):
    line = lines[index]
    if line.startswith("```"):
        code = not code
    elif line.startswith("|"):
        rows = []
        while index < len(lines) and lines[index].startswith("|"):
            row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
            if not all(re.fullmatch(r"[:\- ]+", cell) for cell in row):
                rows.append(row)
            index += 1
        table = document.add_table(rows=0, cols=len(rows[0]))
        table.style = "Light Shading Accent 1"
        for row in rows:
            for cell, value in zip(table.add_row().cells, row):
                cell.text = plain(value)
        continue
    elif line.startswith("!["):
        match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        image = folder / match.group(2)
        shape = document.add_picture(str(image), width=Inches(6.6))
        if shape.height > Inches(7.7):
            scale = Inches(7.7) / shape.height
            shape.width = int(shape.width * scale)
            shape.height = Inches(7.7)
        document.add_paragraph(match.group(1), style="Caption")
    elif line.startswith("#"):
        level = len(line) - len(line.lstrip("#"))
        document.add_heading(plain(line.lstrip("# ")), level=min(level, 3))
    elif line.strip():
        paragraph = document.add_paragraph(plain(line[2:] if line.startswith("- ") else line),
            style="List Bullet" if line.startswith("- ") else None)
        if code:
            for run in paragraph.runs:
                run.font.name = "Consolas"
                run.font.size = Pt(8)
    index += 1
document.save(str(source.with_suffix(".docx")))
set_status("R16", "CHO_KIEM_THU", "report/BaoCao_Local_521H0092.md/.pdf/.docx: đủ 5 chương cho local; chờ EC2 và MSSV thành viên hai")
print("Built local report: Markdown, HTML, PDF and DOCX; EC2 remains pending.")
